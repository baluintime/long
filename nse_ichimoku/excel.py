"""Excel workbook export.

The workbook is the complete record of a scan: a summary sheet, one sheet
per Ichimoku position, and an ``All`` sheet holding every classified stock,
each row carrying both the raw levels and the percentage gap between the
close and every line.
"""

from __future__ import annotations

import logging
import os
import tempfile
from datetime import datetime
from pathlib import Path

import pandas as pd

from .ichimoku import Position
from .report import (
    DETAIL_COLUMNS,
    PERCENT_COLUMNS,
    PRICE_COLUMNS,
    all_readings,
    readings_to_frame,
)
from .screener import IST, ScanResult

log = logging.getLogger(__name__)

SUMMARY_SHEET = "Summary"

#: Every .xlsx is a ZIP archive; Excel rejects anything else outright.
ZIP_MAGIC = b"PK"
PRICE_FORMAT = "#,##0.00"
PERCENT_FORMAT = '0.00"%"'
MAX_COLUMN_WIDTH = 22

HEADER_FILL = "FF1F3864"  # deep blue
HEADER_FONT = "FFFFFFFF"
ABOVE_FILL = "FFE2EFDA"  # light green
BELOW_FILL = "FFFCE4E4"  # light red
MIXED_FILL = "FFFFF2CC"  # light amber

_POSITION_FILLS = {
    Position.ABOVE.value: ABOVE_FILL,
    Position.BELOW.value: BELOW_FILL,
    Position.MIXED.value: MIXED_FILL,
}


def _sorted_frame(readings, by: str | None, ascending: bool) -> pd.DataFrame:
    frame = readings_to_frame(readings)
    if by and not frame.empty:
        frame = frame.sort_values(by, ascending=ascending, kind="stable").reset_index(
            drop=True
        )
    return frame


SKIPPED_COLUMNS = ["Symbol", "Reason"]


def build_skipped_frame(result: ScanResult) -> pd.DataFrame:
    """Every symbol that could not be classified, and why."""
    return pd.DataFrame(
        sorted(result.skipped.items()), columns=SKIPPED_COLUMNS
    )


def build_sheets(result: ScanResult) -> dict[str, pd.DataFrame]:
    """Sheet name -> rows, ordered so the strongest positions come first."""
    return {
        # Widest cushion above the lines first.
        "Above": _sorted_frame(result.above, "% to Nearest Line", ascending=False),
        # Deepest below the lines first.
        "Below": _sorted_frame(result.below, "% to Nearest Line", ascending=True),
        "Mixed": _sorted_frame(result.mixed, "Symbol", ascending=True),
        "All": _sorted_frame(all_readings(result), "Symbol", ascending=True),
        # Named, not just counted — the workbook accounts for every symbol.
        "Skipped": build_skipped_frame(result),
    }


def build_summary(result: ScanResult, universe_size: int, source: str) -> pd.DataFrame:
    as_of = result.as_of
    rows = [
        ("Session (as of)", as_of.date().isoformat() if as_of is not None else "n/a"),
        ("Generated", datetime.now(IST).strftime("%Y-%m-%d %H:%M IST")),
        ("Universe size", universe_size),
        ("Universe source", source),
        ("Evaluated", len(result.readings)),
        ("Above all Ichimoku lines", len(result.above)),
        ("Below all Ichimoku lines", len(result.below)),
        ("Mixed", len(result.mixed)),
        ("Skipped (no or short data)", len(result.skipped)),
        ("Lagging the session", len(result.stale)),
    ]
    return pd.DataFrame(rows, columns=["Metric", "Value"])


def _style_sheet(worksheet, frame: pd.DataFrame, freeze_first_column: bool) -> None:
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    header_fill = PatternFill("solid", fgColor=HEADER_FILL)
    header_font = Font(bold=True, color=HEADER_FONT)

    for cell in worksheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    if frame.empty:
        worksheet.freeze_panes = "A2"
        return

    columns = list(frame.columns)
    for index, name in enumerate(columns, start=1):
        letter = get_column_letter(index)
        if name in PERCENT_COLUMNS:
            number_format = PERCENT_FORMAT
        elif name in PRICE_COLUMNS or name == "Value":
            number_format = PRICE_FORMAT
        else:
            number_format = None
        if number_format:
            for cell in worksheet[letter][1:]:
                cell.number_format = number_format

        longest = max((len(str(v)) for v in frame[name].head(200)), default=0)
        worksheet.column_dimensions[letter].width = min(
            max(len(str(name)) + 3, longest + 2), MAX_COLUMN_WIDTH
        )

    if "Position" in columns:
        position_index = columns.index("Position") + 1
        for row in range(2, len(frame) + 2):
            cell = worksheet.cell(row=row, column=position_index)
            fill_color = _POSITION_FILLS.get(str(cell.value))
            if fill_color:
                cell.fill = PatternFill("solid", fgColor=fill_color)
            cell.alignment = Alignment(horizontal="center")

    worksheet.freeze_panes = "B2" if freeze_first_column else "A2"
    worksheet.auto_filter.ref = worksheet.dimensions


def verify_workbook(path: str | Path) -> None:
    """Raise unless ``path`` is a workbook Excel will actually open.

    Excel reports "the file format or file extension is not valid" for
    anything that is not a real .xlsx — CSV text under an .xlsx name, or a
    file truncated by a crash mid-write. Catching that here means a bad
    file is never handed to the user in the first place.
    """
    file_path = Path(path)
    if not file_path.exists():
        raise ValueError(f"{file_path} was not created")
    if file_path.stat().st_size == 0:
        raise ValueError(f"{file_path} is empty")

    with open(file_path, "rb") as handle:
        magic = handle.read(2)
    if magic != ZIP_MAGIC:
        raise ValueError(
            f"{file_path} is not a valid .xlsx: it starts with {magic!r} rather "
            "than a ZIP header, so Excel will refuse to open it"
        )

    from openpyxl import load_workbook

    book = load_workbook(file_path, read_only=True)
    try:
        if not book.sheetnames:
            raise ValueError(f"{file_path} contains no sheets")
    finally:
        book.close()


def _write_sheets(
    target: Path, summary: pd.DataFrame, sheets: dict[str, pd.DataFrame]
) -> None:
    with pd.ExcelWriter(target, engine="openpyxl") as writer:
        summary.to_excel(writer, sheet_name=SUMMARY_SHEET, index=False)
        for name, frame in sheets.items():
            # An empty frame still writes its header row, so the sheet
            # exists and the columns are documented.
            frame.to_excel(writer, sheet_name=name, index=False)

        for name, frame in {SUMMARY_SHEET: summary, **sheets}.items():
            try:
                _style_sheet(
                    writer.sheets[name], frame, freeze_first_column=name != SUMMARY_SHEET
                )
            except Exception as exc:
                # Formatting is cosmetic; never lose the data over it.
                log.warning("could not format sheet %r: %s", name, exc)


def write_workbook(
    result: ScanResult, path: str | Path, universe_size: int, source: str
) -> Path:
    """Write the full scan to an .xlsx workbook and return its path.

    The workbook is built in a temporary file alongside the destination,
    verified, and only then moved into place, so an interrupted or failing
    run leaves the previous day's file intact rather than a corrupt one.
    """
    out_path = Path(path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    summary = build_summary(result, universe_size, source)
    sheets = build_sheets(result)

    handle, temp_name = tempfile.mkstemp(
        dir=out_path.parent, prefix=f".{out_path.stem}-", suffix=".xlsx"
    )
    os.close(handle)
    temp_path = Path(temp_name)
    try:
        _write_sheets(temp_path, summary, sheets)
        verify_workbook(temp_path)
        os.replace(temp_path, out_path)
    except BaseException:
        temp_path.unlink(missing_ok=True)
        raise

    return out_path


__all__ = [
    "DETAIL_COLUMNS",
    "build_sheets",
    "build_summary",
    "verify_workbook",
    "write_workbook",
]
